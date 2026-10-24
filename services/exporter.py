from io import BytesIO
from openpyxl import load_workbook
import pandas as pd

def template_headers(path):
    return [c.value for c in load_workbook(path,read_only=True).active[1]]
def export_xlsx(df,template_path):
    wb=load_workbook(template_path); ws=wb.active
    headers=[c.value for c in ws[1]]
    if ws.max_row>1:ws.delete_rows(2,ws.max_row-1)
    work=df.copy()
    if 'incluir' in work:work=work[work['incluir']==True]
    for h in headers:
        if h not in work:work[h]=''
    work=work[headers].fillna('')
    for r_idx,row in enumerate(work.itertuples(index=False,name=None),2):
        for c_idx,val in enumerate(row,1):ws.cell(r_idx,c_idx,val)
    out=BytesIO();wb.save(out);return out.getvalue()
def export_tsv(df,template_path):
    headers=template_headers(template_path);work=df.copy()
    if 'incluir' in work:work=work[work['incluir']==True]
    for h in headers:
        if h not in work:work[h]=''
    return work[headers].fillna('').to_csv(index=False,sep='\t',lineterminator='\n').encode('utf-8-sig')
